from __future__ import annotations

import copy
import json
import re
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection
from openpyxl.worksheet.worksheet import Worksheet

from worklog_core import WorklogError, WorklogStore, _atomic_write
from report_content import attach_evidence, apply_review, fingerprint, strip_management_noise


def _report_range(
    store: WorklogStore,
    start: str | date | None = None,
    end: str | date | None = None,
) -> tuple[date, date]:
    if start and not end:
        week = store.week_context(start)
        return week.start, week.end
    if not start and not end:
        week = store.week_context()
        return week.start, week.end
    start_date = date.fromisoformat(start) if isinstance(start, str) else start
    end_date = date.fromisoformat(end) if isinstance(end, str) else end
    if not start_date or not end_date:
        raise WorklogError("请选择完整的开始和结束日期")
    if start_date > end_date:
        raise WorklogError("开始日期不能晚于结束日期")
    return start_date, end_date


def build_report_preview(
    store: WorklogStore,
    start: str | date | None = None,
    end: str | date | None = None,
    *, use_review: bool = True,
) -> dict[str, Any]:
    range_start, range_end = _report_range(store, start, end)
    projects = {item["project_id"]: item for item in store.list_projects()}
    from weekly_export import WEEKLY_NOISE_MARKERS, prepare_weekly_report_items
    all_nodes = store.list_nodes()
    eligible = [item for item in all_nodes
                if datetime.fromisoformat(str(item["occurred_at"])).date() <= range_end
                and strip_management_noise(item.get("summary", ""))]
    nodes = []
    for item in eligible:
        occurred = datetime.fromisoformat(str(item["occurred_at"])).date()
        if range_start <= occurred <= range_end and item.get("include_in_report", True):
            nodes.append(item)

    current_map: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for item in sorted(nodes, key=lambda value: value.get("occurred_at", "")):
        if item.get("kind") == "计划":
            continue
        project = projects.get(item["project_id"], {})
        entry = current_map.setdefault(item["project_id"], {
            "project_id": item["project_id"],
            "project_name": project.get("name", item.get("project_name", "未命名项目")),
            "major_work": project.get("major_work", item.get("major_work", "其他")),
            "key_point": project.get("key_point", item.get("key_point", "其他支持事项")),
            "points": [],
            "source_node_ids": [],
        })
        summary = item.get("summary", "").strip()
        if summary and summary not in entry["points"]:
            entry["points"].append(summary)
        if summary:
            entry["source_node_ids"].append(item["node_id"])

    plan_map: OrderedDict[str, dict[str, Any]] = OrderedDict()
    latest_by_project: dict[str, dict[str, Any]] = {}
    for item in eligible:
        if item.get("kind") == "计划":
            continue
        latest_by_project.setdefault(item["project_id"], item)
    tasks = {item["task_id"]: item for item in store.list_tasks()}
    latest_by_task: dict[str, dict[str, Any]] = {}
    for item in eligible:
        if item.get("task_id") and item.get("kind") != "计划":
            latest_by_task.setdefault(item["task_id"], item)

    def add_plan(project_id: str, action: str) -> None:
        action = action.strip()
        if not action:
            return
        project = projects.get(project_id, {})
        entry = plan_map.setdefault(project_id, {
            "project_id": project_id,
            "project_name": project.get("name", "未命名项目"),
            "major_work": project.get("major_work", "其他"),
            "key_point": project.get("key_point", "其他支持事项"),
            "points": [],
            "source_action_ids": [],
        })
        if action not in entry["points"]:
            entry["points"].append(action)
            entry["source_action_ids"].append(fingerprint([project_id, action]))

    def task_status_at_end(task: dict[str, Any]) -> str:
        status = task.get("status", "进行中")
        for change in reversed(task.get("status_history") or []):
            if str(change.get("at", ""))[:10] > range_end.isoformat():
                status = change.get("from", status)
        return status

    # A blank project-level latest action must not hide other unfinished tasks.
    for task_id, task in tasks.items():
        if str(task.get("created_at", ""))[:10] > range_end.isoformat():
            continue
        status = task_status_at_end(task)
        if status == "已完成":
            continue
        latest = latest_by_task.get(task_id, {})
        action = str(latest.get("next_action") or "").strip()
        if action:
            if status in {"等待中", "阻塞", "待确认"}:
                action = f"【{status}】{action}"
            add_plan(task["project_id"], action)
        else:
            add_plan(task["project_id"], f"【待明确下一步】{task.get('title', '未命名任务')}")

    for project_id, item in latest_by_project.items():
        # Task actions above are authoritative; completed tasks are not revived.
        if item.get("task_id") in tasks:
            continue
        next_action = item.get("next_action", "").strip()
        if not next_action:
            continue
        project = projects.get(project_id, {})
        add_plan(project_id, next_action)
    for item in sorted(nodes, key=lambda value: value.get("occurred_at", "")):
        if item.get("kind") != "计划":
            continue
        latest = latest_by_task.get(item.get("task_id")) if item.get("task_id") else latest_by_project.get(item["project_id"])
        if latest and latest.get("occurred_at", "") >= item.get("occurred_at", ""):
            continue
        entry = plan_map.setdefault(item["project_id"], {
            "project_id": item["project_id"],
            "project_name": item.get("project_name", "未命名项目"),
            "major_work": item.get("major_work", "其他"),
            "key_point": item.get("key_point", "其他支持事项"),
            "points": [],
            "source_action_ids": [],
        })
        summary = item.get("summary", "").strip()
        if summary and summary not in entry["points"]:
            entry["points"].append(summary)
            entry["source_action_ids"].append(fingerprint([item["project_id"], summary]))

    report = attach_evidence({
        "range": {"start": range_start.isoformat(), "end": range_end.isoformat()},
        "current_items": prepare_weekly_report_items(store, list(current_map.values())),
        "plan_items": prepare_weekly_report_items(store, list(plan_map.values())),
    }, eligible, list(tasks.values()), list(projects.values()))
    saved_review = store.drafts_dir / f"{range_start}_{range_end}.content-review.json"
    if use_review and saved_review.exists():
        try:
            report = apply_review(report, json.loads(saved_review.read_text(encoding="utf-8")))
        except (WorklogError, ValueError, TypeError, KeyError):
            report["review_required"] = True
            report["review_warning"] = "已有整理稿与当前记录不一致，请重新核验"
    return report


def _copy_style(source, target) -> None:
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy.copy(source.protection)


def _style_row(ws: Worksheet, row: int, prototypes: tuple[Any, Any, Any]) -> None:
    for col, prototype in enumerate(prototypes, start=1):
        _copy_style(prototype, ws.cell(row=row, column=col))


def _content_text(item: dict[str, Any]) -> str:
    points = [str(point).strip() for point in item.get("points", []) if str(point).strip()]
    if not points:
        points = ["暂无记录"]
    rendered = [str(item.get("project_name", "未命名项目")).strip()]
    for index, point in enumerate(points, start=1):
        rendered.append(f"{index}.{point}")
    return "\n".join(rendered)


def _section_rows(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if items:
        return items
    return [{
        "project_id": "empty",
        "project_name": "暂无记录",
        "major_work": "其他",
        "key_point": "其他支持事项",
        "points": ["暂无记录"],
    }]


def _write_section(
    ws: Worksheet,
    start_row: int,
    *,
    title: str,
    items: list[dict[str, Any]],
    title_prototypes: tuple[Any, Any, Any],
    header_prototypes: tuple[Any, Any, Any],
    data_prototypes: tuple[Any, Any, Any],
) -> int:
    _style_row(ws, start_row, title_prototypes)
    ws.cell(start_row, 1, title)
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=3)
    ws.row_dimensions[start_row].height = 20.4

    header_row = start_row + 1
    _style_row(ws, header_row, header_prototypes)
    for column, value in enumerate(("主要工作", "关键点", "具体内容"), start=1):
        ws.cell(header_row, column, value)
    ws.row_dimensions[header_row].height = 20

    data_start = header_row + 1
    rows = _section_rows(items)
    for offset, item in enumerate(rows):
        row = data_start + offset
        _style_row(ws, row, data_prototypes)
        ws.cell(row, 1, item.get("major_work", "其他"))
        ws.cell(row, 2, item.get("key_point", "其他支持事项"))
        content = _content_text(item)
        ws.cell(row, 3, content)
        ws.cell(row, 3).alignment = copy.copy(data_prototypes[2].alignment)
        ws.cell(row, 3).alignment = Alignment(
            horizontal=ws.cell(row, 3).alignment.horizontal,
            vertical="center",
            wrap_text=True,
        )
        line_count = content.count("\n") + 1
        visual_lines = line_count + sum(max(0, len(line) // 55) for line in content.splitlines())
        ws.row_dimensions[row].height = min(180, max(31.8, 16 * visual_lines + 10))

    end_row = data_start + len(rows) - 1
    for column, field in ((1, "major_work"), (2, "key_point")):
        group_start = data_start
        previous = rows[0].get(field, "")
        for index in range(1, len(rows) + 1):
            current = rows[index].get(field, "") if index < len(rows) else None
            if current != previous:
                group_end = data_start + index - 1
                if group_end > group_start:
                    ws.merge_cells(start_row=group_start, start_column=column, end_row=group_end, end_column=column)
                group_start = data_start + index
                previous = current
    return end_row


def _revision_path(export_dir: Path, base_name: str) -> Path:
    candidate = export_dir / f"{base_name}.xlsx"
    if not candidate.exists():
        return candidate
    revision = 1
    while True:
        candidate = export_dir / f"{base_name}-修订{revision}.xlsx"
        if not candidate.exists():
            return candidate
        revision += 1


def export_weekly_report(
    store: WorklogStore,
    *,
    template_path: str | Path,
    export_dir: str | Path,
    start: str | date | None = None,
    end: str | date | None = None,
    anchor: str | date | None = None,
    current_items: list[dict[str, Any]],
    plan_items: list[dict[str, Any]],
) -> Path:
    template = Path(template_path)
    if not template.exists():
        raise WorklogError(f"Excel模板不存在：{template}")
    output_dir = Path(export_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    if anchor and not start and not end:
        range_start, range_end = _report_range(store, anchor, None)
    else:
        range_start, range_end = _report_range(store, start, end)

    workbook = load_workbook(template)
    ws = workbook.active
    prototypes = {
        "title": tuple(copy.copy(ws.cell(1, col)) for col in range(1, 4)),
        "header": tuple(copy.copy(ws.cell(2, col)) for col in range(1, 4)),
        "plan": tuple(copy.copy(ws.cell(3, col)) for col in range(1, 4)),
        "summary": tuple(copy.copy(ws.cell(14, col)) for col in range(1, 4)),
    }
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)

    plan_end = _write_section(
        ws,
        1,
        title="下周工作计划",
        items=plan_items,
        title_prototypes=prototypes["title"],
        header_prototypes=prototypes["header"],
        data_prototypes=prototypes["plan"],
    )
    separator = plan_end + 1
    ws.row_dimensions[separator].height = 12
    summary_start = separator + 1
    summary_end = _write_section(
        ws,
        summary_start,
        title="本周完成情况",
        items=current_items,
        title_prototypes=prototypes["title"],
        header_prototypes=prototypes["header"],
        data_prototypes=prototypes["summary"],
    )

    ws.sheet_view.zoomScale = 85
    ws.sheet_view.showGridLines = False
    ws.print_area = f"A1:C{summary_end}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.column_dimensions["A"].width = 29.44140625
    ws.column_dimensions["B"].width = 79.33203125
    ws.column_dimensions["C"].width = 115.77734375

    base_name = f"团队周报（{range_start.month}月{range_start.day}日~{range_end.month}月{range_end.day}日）"
    output = _revision_path(output_dir, base_name)
    workbook.save(output)

    draft_payload = {
        "type": "weekly-report-snapshot",
        "range_start": range_start.isoformat(),
        "range_end": range_end.isoformat(),
        "exported_at": datetime.now().replace(microsecond=0).isoformat(),
        "excel_path": str(output),
        "plan_items": plan_items,
        "current_items": current_items,
    }
    draft = store.drafts_dir / f"{range_start.isoformat()}_{range_end.isoformat()}_{output.stem}.md"
    body = [
        "---",
        yaml.safe_dump({key: value for key, value in draft_payload.items() if key not in {"plan_items", "current_items"}}, allow_unicode=True, sort_keys=False).strip(),
        "---",
        "",
        "# 周报导出快照",
        "",
        "## 下周工作计划",
        "",
        "```json",
        json.dumps(plan_items, ensure_ascii=False, indent=2),
        "```",
        "",
        "## 本周完成情况",
        "",
        "```json",
        json.dumps(current_items, ensure_ascii=False, indent=2),
        "```",
        "",
    ]
    _atomic_write(draft, "\n".join(body))
    return output


# 保留旧实现供历史追溯，所有公开调用统一使用经过验证的新版导出器。
from weekly_export import export_weekly_report  # noqa: E402,F401
