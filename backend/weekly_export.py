from __future__ import annotations

import json
import math
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

import xlsxwriter
import yaml
from openpyxl import load_workbook

from worklog_core import WorklogError, WorklogStore, _atomic_write
from report_content import leader_text, PRIVATE_TEXT, apply_review, fingerprint


WEEKLY_EXCLUDED_PROJECT_IDS = set()
WEEKLY_MAJOR_ORDER = {"项目推进": 0, "培训工作": 1, "其他": 2}
WEEKLY_KEY_POINTS = {
    "项目推进": "项目推进",
    "培训工作": "培训工作",
    "其他": "其他支持事项",
}
WEEKLY_NOISE_MARKERS = ("重新分类", "归入培训实施", "项目分类调整", "项目归类调整")


def _report_range(
    store: WorklogStore,
    start: str | date | None = None,
    end: str | date | None = None,
) -> tuple[date, date]:
    if start and not end:
        week = store.week_context(start)
        return week.start, week.end
    if not start and not end:
        week = store.week_context()
        return week.start, week.end
    start_date = date.fromisoformat(start) if isinstance(start, str) else start
    end_date = date.fromisoformat(end) if isinstance(end, str) else end
    if not start_date or not end_date:
        raise WorklogError("请选择完整的开始和结束日期")
    if start_date > end_date:
        raise WorklogError("开始日期不能晚于结束日期")
    return start_date, end_date


def _revision_path(export_dir: Path, base_name: str) -> Path:
    candidate = export_dir / f"{base_name}.xlsx"
    if not candidate.exists():
        return candidate
    revision = 1
    while True:
        candidate = export_dir / f"{base_name}-修订{revision}.xlsx"
        if not candidate.exists():
            return candidate
        revision += 1


def prepare_weekly_report_items(
    store: WorklogStore,
    items: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Apply weekly presentation rules without changing Obsidian history."""
    projects = {item["project_id"]: item for item in store.list_projects()}
    prepared: list[dict[str, Any]] = []
    for item in items:
        project_id = str(item.get("project_id") or "")
        project = projects.get(project_id, {})
        if project_id in WEEKLY_EXCLUDED_PROJECT_IDS:
            continue
        if project.get("display_in_views", True) is False:
            continue
        category_text = " ".join(
            str(project.get(key) or item.get(key) or "")
            for key in ("major_work", "work_domain", "work_category")
        )
        if "个人" in category_text:
            continue

        valid_points = [
            leader_text(str(point))
            for point in item.get("points", [])
            if str(point).strip()
        ]
        if not valid_points:
            continue

        raw_major = str(project.get("major_work") or item.get("major_work") or "其他")
        domain = str(project.get("work_domain") or "")
        if raw_major == "项目推进" or domain == "项目管理":
            major_work = "项目推进"
        elif raw_major == "培训工作" or domain == "培训工作":
            major_work = "培训工作"
        else:
            major_work = "其他"
        prepared.append({
            "project_id": project_id,
            "project_name": str(project.get("name") or item.get("project_name") or "未命名项目"),
            "major_work": major_work,
            "key_point": WEEKLY_KEY_POINTS[major_work],
            "points": list(dict.fromkeys(p for p in valid_points if p)),
            **({"source_node_ids": item["source_node_ids"]} if "source_node_ids" in item else {}),
            **({"source_action_ids": item["source_action_ids"]} if "source_action_ids" in item else {}),
        })
    prepared.sort(key=lambda item: (
        WEEKLY_MAJOR_ORDER.get(item["major_work"], 99),
        item["project_name"],
    ))
    return prepared


def _formats(workbook: xlsxwriter.Workbook) -> dict[str, Any]:
    border = {"border": 1, "border_color": "#000000"}
    return {
        "title": workbook.add_format({
            "font_name": "等线", "font_size": 16, "bold": True,
            "font_color": "#FFFFFF", "bg_color": "#1F4E78",
            "align": "center", "valign": "vcenter", **border,
        }),
        "header": workbook.add_format({
            "font_name": "等线", "font_size": 14, "bold": True,
            "font_color": "#000000", "bg_color": "#8FAADC",
            "align": "center", "valign": "vcenter", **border,
        }),
        "group": workbook.add_format({
            "font_name": "等线", "font_size": 12, "bg_color": "#D9E1F2",
            "align": "center", "valign": "vcenter", "text_wrap": True, **border,
        }),
        "content": workbook.add_format({
            "font_name": "等线", "font_size": 11, "align": "left",
            "valign": "vcenter", "text_wrap": True, **border,
        }),
        "project": workbook.add_format({
            "font_name": "等线", "font_size": 11, "bold": True,
        }),
        "progress": workbook.add_format({
            "font_name": "等线", "font_size": 11,
        }),
    }


def content_text(item: dict[str, Any]) -> str:
    return str(item.get("project_name") or "暂无记录") + "\n" + "\n".join(
        f"{i}.{p}" for i, p in enumerate(item.get("points") or ["暂无记录"], 1))


def content_height(text: str) -> float:
    # CJK glyphs occupy about twice the width of Latin glyphs. Leave wrapping margin.
    lines = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in line) / 105)) for line in text.splitlines())
    return max(54, lines * 17 + 12)


def _write_section(
    worksheet: Any,
    row: int,
    title: str,
    items: list[dict[str, Any]],
    formats: dict[str, Any],
) -> int:
    for column in range(3):
        worksheet.write(row, column, title if column == 0 else "", formats["title"])
    worksheet.set_row(row, 32)
    row += 1
    for column, value in enumerate(("主要工作", "关键点", "具体内容")):
        worksheet.write(row, column, value, formats["header"])
    worksheet.set_row(row, 24)
    row += 1

    rows = items or [{
        "project_name": "暂无记录",
        "major_work": "其他",
        "key_point": "其他支持事项",
        "points": ["暂无记录"],
    }]
    data_start = row
    for item in rows:
        project_name = str(item.get("project_name") or "未命名项目")
        text = content_text(item)
        height = content_height(text)
        if height > 409:
            raise WorklogError(f"{project_name}正文超过单行可展示长度，请按独立事项精简后导出")
        worksheet.write_rich_string(
            row,
            2,
            formats["project"],
            project_name,
            formats["progress"],
            text[len(project_name):],
            formats["content"],
        )
        worksheet.set_row(row, height)
        row += 1

    for column, field in ((0, "major_work"), (1, "key_point")):
        group_start = data_start
        previous = rows[0].get(field, "")
        for index in range(1, len(rows) + 1):
            current = rows[index].get(field, "") if index < len(rows) else None
            if current != previous:
                group_end = data_start + index - 1
                if group_end > group_start:
                    worksheet.merge_range(
                        group_start, column, group_end, column, previous, formats["group"]
                    )
                else:
                    worksheet.write(group_start, column, previous, formats["group"])
                group_start = data_start + index
                previous = current
    return row - 1


def verify_weekly_report(path: str | Path, *, expected_current=None, expected_plan=None) -> dict[str, Any]:
    """Re-open the workbook and verify the content and style contract."""
    workbook = load_workbook(path, data_only=False)
    ws = workbook["Sheet1"]
    first_column = [ws.cell(row, 1).value for row in range(1, ws.max_row + 1)]
    if first_column.count("下周工作计划") != 1 or first_column.count("本周完成情况") != 1:
        raise WorklogError("周报验证失败：两个分区不完整")
    summary_row = first_column.index("本周完成情况") + 1
    title_rows = {1, summary_row}
    for area in ws.merged_cells.ranges:
        if area.min_row in title_rows:
            raise WorklogError("周报验证失败：分区标题行不应合并")
    expected_colors = {"FF1F4E78", "1F4E78"}
    for row in title_rows:
        for column in range(1, 4):
            color = ws.cell(row, column).fill.fgColor.rgb
            if color not in expected_colors:
                raise WorklogError("周报验证失败：分区标题颜色不正确")
    # XlsxWriter stores character widths in Excel internal units; openpyxl reads converted values.
    expected_widths = {"A": 30.140625, "B": 80.0, "C": 116.42578125}
    for column, expected in expected_widths.items():
        actual = float(ws.column_dimensions[column].width or 0)
        if abs(actual - expected) > 0.2:
            raise WorklogError(f"周报验证失败：{column}列宽度不正确")
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, 3).value
        if isinstance(value, str) and "\n" in value:
            if PRIVATE_TEXT.search(value):
                raise WorklogError("周报正文仍含路径、内部编号或邮箱")
            if float(ws.row_dimensions[row].height or 0) < content_height(value):
                raise WorklogError("周报正文行高不足，可能截断")
    for expected, start_row, end_row in ((expected_plan, 3, summary_row - 2), (expected_current, summary_row + 2, ws.max_row)):
        if expected is not None:
            actual = [ws.cell(r, 3).value for r in range(start_row, end_row + 1)]
            wanted = [content_text(i) for i in expected] or ["暂无记录\n1.暂无记录"]
            if actual != wanted:
                raise WorklogError("周报正文与核验稿不一致：存在遗漏、重复或改写")
    result = {"ok": True, "rows": ws.max_row, "summary_row": summary_row,
              "content_checked": expected_current is not None and expected_plan is not None}
    workbook.close()
    return result


def _export_weekly_report(
    store: WorklogStore,
    *,
    template_path: str | Path | None = None,
    export_dir: str | Path,
    start: str | date | None = None,
    end: str | date | None = None,
    anchor: str | date | None = None,
    current_items: list[dict[str, Any]],
    plan_items: list[dict[str, Any]],
    review: dict[str, Any] | None = None,
    source_digest: str | None = None,
) -> Path:
    """Generate and verify the approved three-column weekly report."""
    if anchor and not start and not end:
        range_start, range_end = _report_range(store, anchor, None)
    else:
        range_start, range_end = _report_range(store, start, end)
    from reporting import build_report_preview
    fresh = build_report_preview(store, range_start, range_end, use_review=review is None)
    if source_digest is not None and source_digest != fresh["source_digest"]:
        raise WorklogError("记录或日期范围已变化，请重新加载周报预览")
    if review is not None:
        fresh = apply_review(fresh, review)
    else:
        if fresh["review_required"]:
            raise WorklogError("周报内容尚未核验，请依据当周来源完成事项合并和正文整理后导出（CLI 使用 --review-file）")
        if current_items != fresh["current_items"] or plan_items != fresh["plan_items"]:
            raise WorklogError("导出内容与最新预览不一致；改写或排除事项需提供带来源的整理稿")
    prepared_current = prepare_weekly_report_items(store, fresh["current_items"])
    prepared_plan = prepare_weekly_report_items(store, fresh["plan_items"])
    # Validate text/size before creating any file.
    for item in prepared_current + prepared_plan:
        if PRIVATE_TEXT.search(content_text(item)) or content_height(content_text(item)) > 409:
            raise WorklogError("正文仍含内部信息或过长，请先整理后导出")

    output_dir = Path(export_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    base_name = (
        f"团队周报（{range_start.month}月{range_start.day}日"
        f"~{range_end.month}月{range_end.day}日）"
    )
    output = _revision_path(output_dir, base_name)
    workbook = xlsxwriter.Workbook(output)
    worksheet = workbook.add_worksheet("Sheet1")
    worksheet.hide_gridlines(2)
    worksheet.set_zoom(85)
    worksheet.set_column(0, 0, 29.44)
    worksheet.set_column(1, 1, 79.33)
    worksheet.set_column(2, 2, 115.78)
    worksheet.set_portrait()
    worksheet.fit_to_pages(1, 0)
    formats = _formats(workbook)

    plan_end = _write_section(
        worksheet, 0, "下周工作计划", prepared_plan, formats
    )
    separator = plan_end + 1
    worksheet.set_row(separator, 12)
    summary_start = separator + 1
    summary_end = _write_section(
        worksheet, summary_start, "本周完成情况", prepared_current, formats
    )
    worksheet.print_area(0, 0, summary_end, 2)
    workbook.close()
    verification = verify_weekly_report(output, expected_current=prepared_current, expected_plan=prepared_plan)

    draft_payload = {
        "type": "weekly-report-snapshot",
        "range_start": range_start.isoformat(),
        "range_end": range_end.isoformat(),
        "exported_at": datetime.now().replace(microsecond=0).isoformat(),
        "excel_path": str(output),
        "plan_items": prepared_plan,
        "current_items": prepared_current,
        "source_digest": fresh["source_digest"],
        "content_digest": fingerprint([prepared_current, prepared_plan]),
        "verification": verification,
    }
    draft = store.drafts_dir / (
        f"{range_start.isoformat()}_{range_end.isoformat()}_{output.stem}.md"
    )
    body = [
        "---",
        yaml.safe_dump(
            {
                key: value
                for key, value in draft_payload.items()
                if key not in {"plan_items", "current_items"}
            },
            allow_unicode=True,
            sort_keys=False,
        ).strip(),
        "---",
        "",
        "# 周报导出快照",
        "",
        "## 下周工作计划",
        "",
        "~~~json",
        json.dumps(prepared_plan, ensure_ascii=False, indent=2),
        "~~~",
        "",
        "## 本周完成情况",
        "",
        "~~~json",
        json.dumps(prepared_current, ensure_ascii=False, indent=2),
        "~~~",
        "",
    ]
    _atomic_write(draft, "\n".join(body))
    _atomic_write(output.with_suffix(".verification.json"), json.dumps({
        "range": fresh["range"], "source_digest": fresh["source_digest"],
        "current_items": prepared_current, "plan_items": prepared_plan,
        "verification": verification, "review": fresh.get("review"),
    }, ensure_ascii=False, indent=2))
    if review is not None:
        _atomic_write(store.drafts_dir / f"{range_start}_{range_end}.content-review.json", json.dumps(review, ensure_ascii=False, indent=2))
    return output


def export_weekly_report(store: WorklogStore, **kwargs: Any) -> Path:
    # Revalidate sources and choose a non-overwriting filename under the same Vault lock.
    with store.mutation_lock():
        return _export_weekly_report(store, **kwargs)
